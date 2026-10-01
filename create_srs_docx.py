from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.section import WD_SECTION
from pathlib import Path

OUT = Path(__file__).parent / "GridWise_Requirements_and_Test_Cases.docx"
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.7); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.72); sec.right_margin = Inches(.72)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(9.5)
styles['Normal']._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos'); styles['Normal']._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos')
for name, size in [('Title', 25), ('Heading 1', 16), ('Heading 2', 12)]:
    st=styles[name]; st.font.name='Aptos Display'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=RGBColor(0,0,0)
    st._element.rPr.rFonts.set(qn('w:ascii'), 'Aptos Display'); st._element.rPr.rFonts.set(qn('w:hAnsi'), 'Aptos Display')
    st.paragraph_format.space_before=Pt(12); st.paragraph_format.space_after=Pt(6)

def shade(cell, color):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),color); tcPr.append(shd)
def border_table(table):
    tblPr=table._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        tag=OxmlElement('w:'+side); tag.set(qn('w:val'),'single'); tag.set(qn('w:sz'),'4'); tag.set(qn('w:color'),'D9D9D9'); borders.append(tag)
    tblPr.append(borders)
def add_table(headers, rows, widths=None):
    table=doc.add_table(rows=1, cols=len(headers)); table.alignment=WD_TABLE_ALIGNMENT.CENTER; table.style='Table Grid'; border_table(table)
    for i,h in enumerate(headers):
        c=table.rows[0].cells[i]; c.text=h; shade(c,'173F3A'); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for r in c.paragraphs[0].runs: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
    for ri,row in enumerate(rows):
        cells=table.add_row().cells
        for i,value in enumerate(row):
            cells[i].text=str(value); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if ri%2: shade(cells[i],'F4F8F5')
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
                for r in p.runs: r.font.size=Pt(8)
    if widths:
        for row in table.rows:
            for c,w in zip(row.cells,widths): c.width=Inches(w)
    doc.add_paragraph().paragraph_format.space_after=Pt(3)

title=doc.add_paragraph(style='Title'); title.alignment=WD_ALIGN_PARAGRAPH.LEFT; title.add_run('GridWise IoT Energy Management System')
sub=doc.add_paragraph('Software requirements and test cases'); sub.runs[0].font.size=Pt(13); sub.runs[0].font.color.rgb=RGBColor(77,112,101)
p=doc.add_paragraph('This document defines the functional scope, quality targets, multi-agent workflow, and verification scenarios for the GridWise IoT energy management platform. It is intended for facility operations, engineering, and project stakeholders.'); p.paragraph_format.space_after=Pt(10)

doc.add_heading('1 Purpose and scope', level=1)
doc.add_paragraph('GridWise helps facility operators reduce energy cost and peak demand by combining IoT telemetry with specialised AI agents. The system collects device data, identifies operational opportunities, detects equipment anomalies, and presents safe actions for operator approval or policy-controlled execution.')
doc.add_heading('2 Stakeholders', level=1)
add_table(['Role','Primary need'],[
['Energy manager','Monitor live performance and approve cost-saving actions.'],['Facility technician','Investigate device and equipment alerts.'],['Finance team','Track cost and savings against forecast.'],['System administrator','Configure devices, tariffs, permissions and policies.']], [1.55,5.1])
doc.add_heading('3 Functional requirements', level=1)
add_table(['ID','Requirement','Priority','Acceptance criterion'],[
['FR-01','Ingest timestamped power, energy, voltage, battery and solar telemetry.','Must','Valid readings include device ID and timestamp.'],
['FR-02','Show current demand, daily consumption, daily cost and solar generation.','Must','Values refresh for the selected site.'],
['FR-03','Display a configurable demand trend.','Must','Operator can select day, week or month aggregation.'],
['FR-04','Forecast near-term site demand.','Must','Agent returns a peak-window estimate and confidence.'],
['FR-05','Optimise demand against tariff windows and demand limits.','Must','Agent proposes action window and estimated saving.'],
['FR-06','Detect abnormal device or equipment consumption.','Must','Alert includes severity, evidence and source device.'],
['FR-07','Route outputs through orchestration policy.','Must','Conflicts are ranked and only permitted actions proceed.'],
['FR-08','Allow operator to apply or dismiss a recommendation.','Must','Decision, actor and timestamp are recorded.'],
['FR-09','Show live grid, solar and battery supply mix.','Should','Sources reconcile to demand within tolerance.'],
['FR-10','Maintain an audit trail for decisions and device commands.','Must','Record contains agent, input, action and result.']], [.55,2.7,.62,2.75])
doc.add_heading('4 Non functional requirements', level=1)
add_table(['ID','Requirement','Target'],[
['NFR-01','Dashboard availability','99.5% monthly availability excluding planned maintenance.'],['NFR-02','Telemetry freshness','No older than 60 seconds under normal connectivity.'],['NFR-03','Dashboard responsiveness','Primary views render within 2 seconds.'],['NFR-04','Security','TLS, encrypted credentials, RBAC and auditable commands.'],['NFR-05','Scalability','At least 10,000 devices at one reading per minute.'],['NFR-06','Safety','Policy checks and a reversible command path.']], [.65,2.0,4.0])
doc.add_heading('5 Multi agent workflow', level=1)
for label,txt in [('Telemetry monitor','Validates data quality, late readings and device health.'),('Load forecaster','Predicts demand using historical load, weather and occupancy signals.'),('Tariff optimiser','Compares forecast load with tariff and contracted demand constraints.'),('Anomaly detector','Identifies deviations from equipment baselines.'),('Orchestrator','Scores, deduplicates and applies policy before surfacing a recommendation or command.')]:
    p=doc.add_paragraph(style='Normal'); p.add_run(label+'. ').bold=True; p.add_run(txt)
doc.add_heading('6 Test cases', level=1)
add_table(['ID','Scenario','Steps','Expected result'],[
['TC-01','Dashboard loads','Open the dashboard.','KPIs, charts, agent panel and alerts are visible.'],
['TC-02','Change time range','Select Week then Month.','Selected range is highlighted and feedback confirms view.'],
['TC-03','Run agent scan','Select Run agent scan.','Scanning state completes and outcome is confirmed.'],
['TC-04','Apply recommendation','Select Apply recommendation.','Success message appears; action is marked applied.'],
['TC-05','Dismiss recommendation','Refresh and select Dismiss.','Recommendation disappears and dismissal is confirmed.'],
['TC-06','Telemetry validation','Send missing ID or invalid numeric reading.','Reading is rejected or quarantined with validation record.'],
['TC-07','Peak forecast','Provide rising demand before a tariff peak.','Forecast and a lower-cost action are produced.'],
['TC-08','Abnormal equipment load','Send AHU power 18% over baseline for 47 minutes.','Alert records device, deviation and severity.'],
['TC-09','Policy enforcement','Submit command outside allowed equipment hours.','Orchestrator blocks action and records reason.'],
['TC-10','Audit record','Apply an approved command.','Audit includes operator, recommendation, command status and timestamp.'],
['TC-11','Responsive layout','View at 375 px and 1024 px widths.','No horizontal overflow; content remains legible.'],
['TC-12','Access control','Attempt a command as read-only user.','System blocks the command with authorisation message.']], [.55,1.25,2.35,2.3])
doc.add_heading('7 Suggested production architecture', level=1)
doc.add_paragraph('Use MQTT or HTTPS gateways for device integration, a time-series database and event queue for telemetry, containerised agent services with a policy engine, an authenticated API, a responsive web dashboard with WebSocket updates, and central logs with immutable audit storage.')
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; footer.add_run('GridWise IoT Energy Management System  |  Requirements and Test Cases').font.size=Pt(8)
doc.save(OUT)
print(OUT)
